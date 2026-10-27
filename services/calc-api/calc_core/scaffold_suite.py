"""Generate the complete scaffold-calculation regression DOCX suite."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
from datetime import datetime
from pathlib import Path
from typing import Any, Sequence

from docx import Document

from . import modules as _builtin_modules  # noqa: F401  (registers built-ins)
from . import registry


SCAFFOLD_MODULE_KEYS = (
    "scaffold-fastener",
    "scaffold-fastener-pinming",
    "scaffold-tie-member",
    "cantilever-scaffold",
    "full-hall-scaffold",
    "suspended-basket",
    "door-scaffold",
    "attached-lifting-scaffold",
    "ground-operation-platform",
    "cantilever-operation-platform",
)

SUITE_SCOPE = (
    "PmStudio 自算回归，不是品茗 oracle；本套件未启动、未操控品茗软件，"
    "不得据此声称完成品茗对照或授权操作。"
)

REPO_ROOT = Path(__file__).resolve().parents[3]
DEFAULT_OUTPUT_DIR = REPO_ROOT / "oracle" / "platform-suite-20261027"
DEFAULT_QA_PATH = REPO_ROOT / "docs" / "qa" / "2026-10-27-scaffold-suite.md"


def _safe_filename_part(value: str) -> str:
    invalid = '<>:"/\\|?*'
    cleaned = "".join("_" if character in invalid else character for character in value)
    return cleaned.strip(" .") or "module"


def _result_table(doc: Document):
    matches = [
        table
        for table in doc.tables
        if table.rows and table.cell(0, 0).text.strip() == "项次"
    ]
    if len(matches) != 1:
        raise RuntimeError(f"expected one result table, found {len(matches)}")
    return matches[0]


def _failure_record(result: dict[str, Any]) -> dict[str, Any]:
    return {
        "no": result["no"],
        "key": result["key"],
        "item": result["item"],
        "value": result["value"],
        "limit": result["limit"],
        "unit": result["unit"],
        "control": bool(result.get("control", False)),
        "clause": result.get("clause", ""),
    }


def _actual_summary(payload: dict[str, Any]) -> dict[str, int]:
    results = payload.get("results", [])
    return {
        **{
            key: value
            for key, value in (payload.get("summary") or {}).items()
            if key not in {"total", "failed", "control_failed"}
        },
        "total": len(results),
        "failed": sum(1 for result in results if not result.get("ok", False)),
        "control_failed": sum(
            1
            for result in results
            if not result.get("ok", False) and result.get("control", False)
        ),
    }


def _write_text_atomic(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.tmp")
    temporary.write_text(text, encoding="utf-8", newline="\n")
    temporary.replace(path)


def _qa_markdown(manifest: dict[str, Any]) -> str:
    lines = [
        "# 2026-10-27 脚手架模块 DOCX 全量回归",
        "",
        f"> {manifest['scope']}",
        "",
        "本记录覆盖当前注册的 10 个脚手架模块，均以 `run({})` 的默认参数调用已注册 "
        "`ModuleDef.run`，再以包含正确 `moduleKey`、`moduleTitle` 和 "
        "`documentTitle` 的 `meta` 调用 `ModuleDef.report`。",
        "",
        "| 模块数 | 结果总数 | 失败总数 | 控制项失败总数 |",
        "| ---: | ---: | ---: | ---: |",
        (
            f"| {manifest['module_count']} | {manifest['result_count']} | "
            f"{manifest['failed_count']} | {manifest['control_failed_count']} |"
        ),
        "",
    ]

    for entry in manifest["modules"]:
        report_link = (
            f"../../oracle/platform-suite-20261027/{entry['报告文件']}"
        )
        lines.extend(
            [
                f"## {entry['sequence']:02d}. {entry['title']}",
                "",
                f"- 模块 key：`{entry['key']}`",
                f"- 标准：`{entry['standard']}`",
                f"- verified：`{str(entry['verified']).lower()}`",
                f"- 分类：{entry['category']}",
                (
                    f"- 汇总：结果 {entry['结果数']} 项，失败 "
                    f"{entry['summary']['failed']} 项，控制项失败 "
                    f"{entry['summary']['control_failed']} 项"
                ),
                (
                    f"- 报告文件：[{entry['报告文件']}]({report_link})，"
                    f"{entry['DOCX大小']} bytes，SHA256 `{entry['SHA256']}`，"
                    f"内嵌图片 {entry['内嵌图片数']} 张，结果表 "
                    f"{entry['结果表行数']} 行（含表头）"
                ),
                "",
                f"**通过项（{len(entry['通过项'])}）**："
                + ("；".join(entry["通过项"]) if entry["通过项"] else "无"),
                "",
            ]
        )
        if entry["失败项"]:
            lines.append(f"**失败项（{len(entry['失败项'])}）**：")
            lines.append("")
            for failure in entry["失败项"]:
                control = "是" if failure["control"] else "否"
                lines.append(
                    f"- 第 {failure['no']} 项 {failure['item']}："
                    f"{failure['value']} {failure['unit']} / 限值 "
                    f"{failure['limit']} {failure['unit']}；控制项={control}；"
                    f"`{failure['key']}`；{failure['clause']}"
                )
        else:
            lines.append("**失败项（0）**：无")
        lines.append("")

    lines.extend(
        [
            "## 边界声明",
            "",
            "- 本套件只验证 PmStudio 自身默认参数、模块注册、DOCX 生成、图表内嵌、"
            "结果表与失败项渲染的可重复性。",
            "- 所有模块仍为 `verified=false`。本套件不替代规范复核、工程审查、"
            "品茗同参数逐项对照或正式签章。",
            "- 品茗 oracle 的独立证据只来自仓库中另行列明来源的导出与结构化记录；"
            "本次未操控品茗，也未绕过任何授权。",
            "",
        ]
    )
    return "\n".join(lines)


def generate_scaffold_suite(
    output_dir: str | Path | None = None,
    qa_path: str | Path | None = None,
    *,
    generated_at: datetime | None = None,
) -> dict[str, Any]:
    """Generate ten DOCX reports plus manifest and QA records."""
    destination = Path(output_dir) if output_dir is not None else DEFAULT_OUTPUT_DIR
    qa_destination = Path(qa_path) if qa_path is not None else DEFAULT_QA_PATH
    destination.mkdir(parents=True, exist_ok=True)
    timestamp = generated_at or datetime.now().astimezone()
    report_date = f"{timestamp:%Y年%m月%d日}"
    entries: list[dict[str, Any]] = []

    for sequence, module_key in enumerate(SCAFFOLD_MODULE_KEYS, start=1):
        definition = registry.require(module_key)
        payload = definition.run({})
        if payload.get("module") != module_key:
            raise RuntimeError(
                f"{module_key} run() returned module={payload.get('module')!r}"
            )

        results = payload.get("results", [])
        passing = [result["item"] for result in results if result.get("ok", False)]
        failed = [
            _failure_record(result)
            for result in results
            if not result.get("ok", False)
        ]
        summary = _actual_summary(payload)
        document_title = f"{definition.title}计算书"
        meta = {
            "moduleKey": definition.key,
            "moduleTitle": definition.title,
            "documentTitle": document_title,
            "reportNo": f"PM-REG-{timestamp:%Y%m%d}-{sequence:03d}",
            "projectName": "PmStudio 脚手架默认参数回归（非品茗 oracle）",
            "clientName": "PmStudio 自算回归",
            "orgName": "PmStudio",
            "date": report_date,
        }
        content = definition.report(payload.get("params") or {}, payload, meta)
        if not content.startswith(b"PK"):
            raise RuntimeError(f"{module_key} report() did not return DOCX bytes")

        doc = Document(io.BytesIO(content))
        result_table = _result_table(doc)
        table_rows = len(result_table.rows)
        if table_rows != len(results) + 1:
            raise RuntimeError(
                f"{module_key} result table has {table_rows} rows for "
                f"{len(results)} results"
            )
        if module_key != "scaffold-tie-member" and not doc.inline_shapes:
            raise RuntimeError(f"{module_key} report contains no inline figures")

        filename = (
            f"{sequence:02d}_{definition.key}_"
            f"{_safe_filename_part(definition.title)}.docx"
        )
        report_path = destination / filename
        report_path.write_bytes(content)
        digest = hashlib.sha256(content).hexdigest()
        entries.append(
            {
                "sequence": sequence,
                "key": definition.key,
                "title": definition.title,
                "standard": str(payload.get("standard") or definition.standard),
                "verified": bool(payload.get("verified", definition.verified)),
                "category": definition.category,
                "summary": summary,
                "通过项": passing,
                "失败项": failed,
                "结果数": len(results),
                "报告文件": filename,
                "DOCX大小": len(content),
                "SHA256": digest,
                "内嵌图片数": len(doc.inline_shapes),
                "结果表行数": table_rows,
            }
        )

    manifest = {
        "schema_version": 1,
        "generated_at": timestamp.isoformat(timespec="seconds"),
        "scope": SUITE_SCOPE,
        "artifact_kind": "pmsstudio-self-calculation-regression",
        "module_count": len(entries),
        "result_count": sum(entry["结果数"] for entry in entries),
        "failed_count": sum(entry["summary"]["failed"] for entry in entries),
        "control_failed_count": sum(
            entry["summary"]["control_failed"] for entry in entries
        ),
        "modules": entries,
    }
    _write_text_atomic(
        destination / "manifest.json",
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
    )
    _write_text_atomic(qa_destination, _qa_markdown(manifest))
    return manifest


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Generate the 10-module PmStudio scaffold DOCX regression suite."
    )
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--qa-path", type=Path, default=DEFAULT_QA_PATH)
    args = parser.parse_args(argv)

    manifest = generate_scaffold_suite(args.output_dir, args.qa_path)
    print(
        f"generated={manifest['module_count']} "
        f"results={manifest['result_count']} "
        f"failed={manifest['failed_count']} "
        f"output={Path(args.output_dir).resolve()}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
