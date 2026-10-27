import io

import pytest
from docx import Document

from calc_core import modules  # noqa: F401  (registers modules)
from calc_core import registry, report
from calc_core.scaffold_report_figures import MODULE_TITLES, build_scaffold_figures


PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"


def test_tie_member_is_in_figure_title_catalog():
    assert MODULE_TITLES["scaffold-tie-member"] == "连墙件计算"


@pytest.mark.parametrize("module_key", MODULE_TITLES)
def test_scaffold_figures_have_titles_and_png_signatures(module_key):
    defn = registry.get(module_key)
    payload = defn.run({})
    figures = build_scaffold_figures(module_key, payload)

    assert 1 <= len(figures) <= 4
    for png, caption in figures:
        assert png.startswith(PNG_SIGNATURE)
        assert defn.title in caption
        assert "简化计算模型" in caption


def test_unsupported_module_has_no_figures():
    assert build_scaffold_figures("scaffold-fastener", {}) == []


def test_pinming_figures_cover_long_transverse_and_both_poles():
    defn = registry.get("scaffold-fastener-pinming")
    payload = defn.run({})
    figures = build_scaffold_figures(defn.key, payload)

    assert len(figures) == 4
    details = [caption for _, caption in figures]
    assert any("纵向水平杆（三跨连续梁）" in text for text in details)
    assert any("横向水平杆（简支梁，P@L/3 + 自重" in text for text in details)
    assert any("单立杆轴力与风弯矩" in text for text in details)
    assert any("双立杆轴力与风弯矩" in text for text in details)
    assert any("单样本系数" in text for text in details)
    assert all(png.startswith(PNG_SIGNATURE) for png, _ in figures)


def test_tie_member_report_has_elevation_and_plan_figures():
    defn = registry.require("scaffold-tie-member")
    payload = defn.run({})
    figures = build_scaffold_figures(defn.key, payload)
    assert len(figures) == 2
    assert "正立面" in figures[0][1]
    assert "平面" in figures[1][1]
    assert all(png.startswith(PNG_SIGNATURE) for png, _ in figures)


def test_tie_member_generic_report_uses_chinese_parameter_labels():
    defn = registry.require("scaffold-tie-member")
    payload = defn.run({})
    content = report.build_generic_docx(
        payload["params"],
        payload,
        {"moduleKey": defn.key, "moduleTitle": defn.title},
    )
    doc = Document(io.BytesIO(content))
    text = "\n".join(
        cell.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
    )

    assert "连墙件竖向间距" in text
    assert "净截面面积" in text
    assert "架体类型" in text
    assert len(doc.inline_shapes) == 2


@pytest.mark.parametrize(
    "module_key",
    [
        "attached-lifting-scaffold",
        "full-hall-scaffold",
        "ground-operation-platform",
        "scaffold-fastener-pinming",
    ],
)
def test_generic_report_embeds_figures_and_keeps_results(module_key):
    defn = registry.get(module_key)
    payload = defn.run({})
    meta = {"moduleKey": module_key, "moduleTitle": defn.title}
    content = report.build_generic_docx(payload["params"], payload, meta)
    doc = Document(io.BytesIO(content))

    assert len(doc.inline_shapes) > 0
    headings = [p.text for p in doc.paragraphs if "Heading 1" in p.style.name]
    assert "三、内力与变形图" in headings
    assert "四、验算结果" in headings
    assert "五、结论" in headings
    assert any("简化计算模型" in p.text for p in doc.paragraphs)

    results_table = next(table for table in doc.tables if table.cell(0, 0).text == "项次")
    assert len(results_table.rows) == len(payload["results"]) + 1
    assert {
        results_table.cell(row, 1).text
        for row in range(1, len(results_table.rows))
    } == {result["item"] for result in payload["results"]}


def test_generic_report_without_figures_keeps_original_numbering():
    defn = registry.get("full-hall-scaffold")
    payload = defn.run({})
    content = report.build_generic_docx(payload["params"], payload, {})
    doc = Document(io.BytesIO(content))
    headings = [p.text for p in doc.paragraphs if "Heading 1" in p.style.name]

    assert len(doc.inline_shapes) == 0
    assert "三、验算结果" in headings
    assert "四、结论" in headings
