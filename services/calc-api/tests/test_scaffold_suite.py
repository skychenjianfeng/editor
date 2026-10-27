import hashlib
import io
import json
import re
from datetime import datetime, timezone

from docx import Document

from calc_core import registry, report
from calc_core.scaffold_suite import (
    SCAFFOLD_MODULE_KEYS,
    generate_scaffold_suite,
)


MANIFEST_FIELDS = {
    "key",
    "title",
    "standard",
    "verified",
    "category",
    "summary",
    "失败项",
    "结果数",
    "DOCX大小",
    "SHA256",
    "内嵌图片数",
    "结果表行数",
}


def test_suite_generates_ten_utf8_named_reports_manifest_and_qa(tmp_path):
    output_dir = tmp_path / "平台套件"
    qa_path = tmp_path / "文档" / "验收.md"
    manifest = generate_scaffold_suite(
        output_dir,
        qa_path,
        generated_at=datetime(2026, 10, 27, 9, 30, tzinfo=timezone.utc),
    )

    assert manifest["module_count"] == 10
    assert manifest["artifact_kind"] == "pmsstudio-self-calculation-regression"
    assert "不是品茗 oracle" in manifest["scope"]
    assert [entry["key"] for entry in manifest["modules"]] == list(SCAFFOLD_MODULE_KEYS)

    report_names = set()
    for entry in manifest["modules"]:
        assert MANIFEST_FIELDS <= entry.keys()
        assert re.fullmatch(
            rf"{entry['sequence']:02d}_{re.escape(entry['key'])}_.+\.docx",
            entry["报告文件"],
        )
        assert entry["报告文件"] not in report_names
        report_names.add(entry["报告文件"])

        report_path = output_dir / entry["报告文件"]
        content = report_path.read_bytes()
        assert content.startswith(b"PK")
        assert entry["DOCX大小"] == len(content)
        assert entry["SHA256"] == hashlib.sha256(content).hexdigest()

        doc = Document(io.BytesIO(content))
        assert entry["内嵌图片数"] == len(doc.inline_shapes)
        assert entry["内嵌图片数"] > 0
        result_table = next(
            table
            for table in doc.tables
            if table.rows and table.cell(0, 0).text.strip() == "项次"
        )
        assert entry["结果表行数"] == len(result_table.rows)
        assert entry["结果表行数"] == entry["结果数"] + 1
        assert entry["summary"]["total"] == entry["结果数"]
        assert entry["summary"]["failed"] == len(entry["失败项"])

    manifest_path = output_dir / "manifest.json"
    parsed_manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert parsed_manifest == manifest

    qa_text = qa_path.read_text(encoding="utf-8")
    assert "不是品茗 oracle" in qa_text
    assert "未启动、未操控品茗软件" in qa_text
    for entry in manifest["modules"]:
        assert f"## {entry['sequence']:02d}. {entry['title']}" in qa_text
        assert entry["报告文件"] in qa_text
        assert f"**通过项（{len(entry['通过项'])}）**" in qa_text
        assert f"**失败项（{len(entry['失败项'])}）**" in qa_text


def test_generic_report_renders_conformity_and_actual_failure_rows():
    definition = registry.require("suspended-basket")
    payload = definition.run({})
    payload["summary"]["failed"] = 0
    payload["results"][0]["ok"] = False

    content = definition.report(
        payload["params"],
        payload,
        {
            "moduleKey": definition.key,
            "moduleTitle": definition.title,
            "documentTitle": f"{definition.title}计算书",
        },
    )
    doc = Document(io.BytesIO(content))
    paragraphs = [paragraph.text for paragraph in doc.paragraphs]

    assert "结果：通过 4/4 项 → 满足要求" in paragraphs
    assert (
        "本次验算共 8 项，其中 1 项不满足：吊篮平台主梁抗弯。"
        in paragraphs
    )

    result_table = next(
        table
        for table in doc.tables
        if table.rows and table.cell(0, 0).text.strip() == "项次"
    )
    assert result_table.cell(1, 5).text == "不满足要求"


def test_generic_report_uses_greater_or_equal_for_strength_factors():
    definition = registry.require("suspended-basket")
    payload = definition.run({})
    stability = next(
        result
        for result in payload["results"]
        if result["key"] == "suspension_stability"
    )

    assert report._result_outcome_text(stability) == (
        f"结果：{stability['value']} {stability['unit']} ≥ 3 → 满足要求"
    )
